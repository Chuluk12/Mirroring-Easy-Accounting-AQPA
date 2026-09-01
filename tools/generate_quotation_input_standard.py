from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Pt


ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "docs" / "Ketetapan_Alur_Input_Quotation.docx"


def add_heading(document, text, level=1):
    heading = document.add_heading(text, level=level)
    heading.paragraph_format.space_before = Pt(8)
    heading.paragraph_format.space_after = Pt(4)
    return heading


def add_bullets(document, rows):
    for row in rows:
        paragraph = document.add_paragraph(style="List Bullet")
        paragraph.add_run(row)


document = Document()
section = document.sections[0]
section.top_margin = Cm(1.8)
section.bottom_margin = Cm(1.8)
section.left_margin = Cm(2)
section.right_margin = Cm(2)

styles = document.styles
styles["Normal"].font.name = "Arial"
styles["Normal"].font.size = Pt(10)
styles["Title"].font.name = "Arial"
styles["Title"].font.size = Pt(18)

title = document.add_paragraph(style="Title")
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
title.add_run("KETETAPAN ALUR INPUT QUOTATION")
subtitle = document.add_paragraph()
subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
subtitle.add_run("PT AQPA INDONESIA\nModul Quotation").bold = True

add_heading(document, "1. Tujuan", 1)
document.add_paragraph(
    "Ketetapan ini menjadi standar pengisian quotation agar harga jual, biaya pendukung, "
    "diskon, pajak, dan margin dihitung secara konsisten serta mudah diperiksa."
)

add_heading(document, "2. Urutan Wajib Pengisian", 1)
steps = [
    ("1", "Data barang", "Wajib diisi lebih dahulu: deskripsi dan spesifikasi, quantity, UOM, serta harga modal barang."),
    ("2", "Biaya pendukung", "Diisi sesuai kebutuhan: CF, MF, instalasi, pengiriman, packaging, dan biaya lainnya."),
    ("3", "Target markup", "Wajib diisi setelah data barang lengkap. Metode dapat berupa persentase atau nominal keuntungan."),
    ("4", "Diskon dan PPN", "Diisi setelah target markup valid. Diskon bersifat opsional; PPN mengikuti ketentuan transaksi."),
    ("5", "Harga jual", "Sistem menghitung Unit Price otomatis. Harga dapat diubah manual setelah mode otomatis dinonaktifkan."),
]
table = document.add_table(rows=1, cols=3)
table.style = "Table Grid"
headers = table.rows[0].cells
headers[0].text = "Tahap"
headers[1].text = "Bagian"
headers[2].text = "Ketetapan"
for number, name, rule in steps:
    cells = table.add_row().cells
    cells[0].text = number
    cells[1].text = name
    cells[2].text = rule

add_heading(document, "3. Prasyarat dan Penguncian Input", 1)
add_bullets(document, [
    "Biaya pendukung dan Target Markup tidak dapat diinput sebelum seluruh data barang lengkap.",
    "Diskon, PPN, pilihan harga otomatis, dan Unit Price tidak dapat digunakan sebelum Target Markup valid.",
    "Unit Price terkunci selama pilihan Hitung harga jual otomatis aktif.",
    "Bagian yang belum dapat digunakan tetap ditampilkan dalam kondisi nonaktif beserta petunjuk data yang harus dilengkapi.",
    "Quotation tidak dapat disimpan apabila data wajib belum lengkap atau nilai tidak memenuhi ketentuan.",
])

add_heading(document, "4. Ketentuan Target Markup", 1)
add_bullets(document, [
    "Target Markup wajib lebih besar dari 0.",
    "Jika metode Persentase digunakan, nilai maksimal adalah 100%.",
    "Jika metode Nominal Keuntungan digunakan, batas 100% tidak berlaku.",
    "Target markup dihitung dari total harga modal barang ditambah seluruh biaya pendukung.",
])

add_heading(document, "5. Ketentuan Biaya Pendukung", 1)
add_bullets(document, [
    "Biaya dapat diinput dalam Rupiah atau persentase.",
    "Biaya CF, MF, instalasi, dan packaging otomatis diterapkan ke seluruh item saat pertama kali diisi.",
    "Marketing dapat mengubah pilihan item yang menerima alokasi melalui checkbox pada tabel barang.",
    "Pengiriman dan biaya lainnya dialokasikan secara proporsional berdasarkan harga modal setiap item.",
    "Total alokasi per item harus sama dengan total biaya pendukung quotation.",
])

add_heading(document, "6. Urutan dan Rumus Perhitungan", 1)
formula_rows = [
    ("Total Modal Barang", "Jumlah (Quantity × Harga Modal Barang)"),
    ("Total Biaya", "Total Modal Barang + Total Biaya Pendukung"),
    ("Target Keuntungan", "Total Biaya × Target Markup % atau Nominal Keuntungan"),
    ("Target Penjualan Bersih", "Total Biaya + Target Keuntungan"),
    ("Harga Sebelum Diskon", "Target Penjualan Bersih disesuaikan agar target tetap tercapai setelah diskon"),
    ("PPN", "Penjualan setelah diskon × Tarif PPN"),
    ("Total Tagihan", "Penjualan setelah diskon + PPN"),
    ("Margin Aktual", "Penjualan setelah diskon − Modal Barang − Biaya Pendukung"),
]
formula_table = document.add_table(rows=1, cols=2)
formula_table.style = "Table Grid"
formula_table.rows[0].cells[0].text = "Komponen"
formula_table.rows[0].cells[1].text = "Rumus"
for component, formula in formula_rows:
    cells = formula_table.add_row().cells
    cells[0].text = component
    cells[1].text = formula

add_heading(document, "7. Ketentuan Diskon dan PPN", 1)
add_bullets(document, [
    "Diskon dapat menggunakan metode persentase atau nominal Rupiah.",
    "Harga rekomendasi memperhitungkan diskon agar target keuntungan tetap tercapai.",
    "PPN ditambahkan pada nilai tagihan customer dan tidak dihitung sebagai keuntungan perusahaan.",
])

add_heading(document, "8. Pemeriksaan Sebelum Menyimpan", 1)
add_bullets(document, [
    "Pastikan data customer dan informasi quotation telah benar.",
    "Pastikan seluruh item memiliki deskripsi, quantity, UOM, dan harga modal barang.",
    "Pastikan biaya pendukung telah dialokasikan ke item yang sesuai.",
    "Periksa Target Penjualan Bersih, Harga Rekomendasi, Total Tagihan, dan Margin Penjualan Rencana.",
    "Jika Unit Price diubah manual, pastikan nilainya tidak lebih rendah dari harga modal barang.",
])

document.add_paragraph()
note = document.add_paragraph()
note.add_run("Catatan: ").bold = True
note.add_run(
    "Sistem menggunakan penguncian bertahap untuk mengurangi kesalahan input. "
    "Marketing mengikuti urutan dari atas ke bawah sampai harga jual terbentuk."
)

OUTPUT.parent.mkdir(parents=True, exist_ok=True)
document.save(OUTPUT)
print(OUTPUT)
